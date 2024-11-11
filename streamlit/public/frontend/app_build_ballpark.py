import openpyxl
from public.backend.globals import *
from snowflake.snowpark.functions import lit, col

size_formula="""=_xlfn.ifna(_xlfn.SWITCH(INVENTORY[[#This Row],[Object_Type]],"TABLE",IF(INVENTORY[[#This Row],[ContentLines]]>500,"LARGE",IF(AND(INVENTORY[[#This Row],[ContentLines]]<=500,INVENTORY[[#This Row],[ContentLines]]>=200),"MEDIUM","SMALL")),"VIEW", IF(INVENTORY[[#This Row],[ContentLines]]>500,"LARGE",IF(AND(INVENTORY[[#This Row],[ContentLines]]<=500,INVENTORY[[#This Row],[ContentLines]]>=200),"MEDIUM","SMALL")),"JOIN INDEX", IF(INVENTORY[[#This Row],[ContentLines]]>500,"LARGE",IF(AND(INVENTORY[[#This Row],[ContentLines]]<=500,INVENTORY[[#This Row],[ContentLines]]>=200),"MEDIUM","SMALL")),"MACRO", IF(INVENTORY[[#This Row],[ContentLines]]>500,"LARGE",IF(AND(INVENTORY[[#This Row],[ContentLines]]<=500,INVENTORY[[#This Row],[ContentLines]]>=200),"MEDIUM","SMALL")),"PROCEDURE", IF(INVENTORY[[#This Row],[ContentLines]]>500,"LARGE",IF(AND(INVENTORY[[#This Row],[ContentLines]]<=500,INVENTORY[[#This Row],[ContentLines]]>=200),"MEDIUM","SMALL")),"FUNCTION", IF(INVENTORY[[#This Row],[ContentLines]]>500,"LARGE",IF(AND(INVENTORY[[#This Row],[ContentLines]]<=500,INVENTORY[[#This Row],[ContentLines]]>=100),"MEDIUM","SMALL")),"BTEQ", IF(INVENTORY[[#This Row],[ContentLines]]>500,"LARGE",IF(AND(INVENTORY[[#This Row],[ContentLines]]<=500,INVENTORY[[#This Row],[ContentLines]]>=150),"MEDIUM","SMALL")),"FASTLOAD", IF(INVENTORY[[#This Row],[ContentLines]]>500,"LARGE",IF(AND(INVENTORY[[#This Row],[ContentLines]]<=500,INVENTORY[[#This Row],[ContentLines]]>=150),"MEDIUM","SMALL")),"MULTILOAD", IF(INVENTORY[[#This Row],[ContentLines]]>500,"LARGE",IF(AND(INVENTORY[[#This Row],[ContentLines]]<=500,INVENTORY[[#This Row],[ContentLines]]>=150),"MEDIUM","SMALL")),"TPT", IF(INVENTORY[[#This Row],[ContentLines]]>500,"LARGE",IF(AND(INVENTORY[[#This Row],[ContentLines]]<=500,INVENTORY[[#This Row],[ContentLines]]>=150),"MEDIUM","SMALL")),"TPUMP", IF(INVENTORY[[#This Row],[ContentLines]]>500,"LARGE",IF(AND(INVENTORY[[#This Row],[ContentLines]]<=500,INVENTORY[[#This Row],[ContentLines]]>=150),"MEDIUM","SMALL")),"PYTHON", IF(INVENTORY[[#This Row],[ContentLines]]>1000,"LARGE",IF(AND(INVENTORY[[#This Row],[ContentLines]]<=1000,INVENTORY[[#This Row],[ContentLines]]>=400),"MEDIUM","SMALL")),"SCALA", IF(INVENTORY[[#This Row],[ContentLines]]>1000,"LARGE",IF(AND(INVENTORY[[#This Row],[ContentLines]]<=1000,INVENTORY[[#This Row],[ContentLines]]>=300),"MEDIUM","SMALL")),"JAVA", IF(INVENTORY[[#This Row],[ContentLines]]>500,"LARGE",IF(AND(INVENTORY[[#This Row],[ContentLines]]<=500,INVENTORY[[#This Row],[ContentLines]]>=150),"MEDIUM","SMALL")),"NOTEBOOK", IF(INVENTORY[[#This Row],[ContentLines]]>500,"LARGE",IF(AND(INVENTORY[[#This Row],[ContentLines]]<=500,INVENTORY[[#This Row],[ContentLines]]>200),"MEDIUM","SMALL")),"SQL", IF(INVENTORY[[#This Row],[ContentLines]]>500,"LARGE",IF(AND(INVENTORY[[#This Row],[ContentLines]]<=500,INVENTORY[[#This Row],[ContentLines]]>200),"MEDIUM","SMALL"))),"N/A")"""
manual_time_formula="""=_xlfn.ifna(VLOOKUP(INVENTORY[[#This Row],[Object_Type]]&INVENTORY[[#This Row],[Size]],MANUAL_TIME[#All],4,FALSE),0)"""
auto_time_formula="""=_xlfn.ifna(VLOOKUP(INVENTORY[[#This Row],[Object_Type]]&INVENTORY[[#This Row],[Size]],AUTOMATED_TIME[#All],4,FALSE),0)"""
added_effort_formula="""=INVENTORY[[#This Row],[DF Aliases]]*0.05+INVENTORY[[#This Row],[RDD Map]]*0.8+INVENTORY[[#This Row],[UDF]]*0.5+INVENTORY[[#This Row],[Spark FS]]*2"""

def getFileExtension(filename):
    root,ext = os.path.splitext(filename)
    ext = ext.replace(".","")
    return ext

def getObjectType(filename):
    obj_type = getFileExtension(filename).replace(".","").upper()
    if obj_type == "PY":
        return "PYTHON"
    elif obj_type in ['IPYNB','DBC','DBX']:
        return "NOTEBOOK"
    return obj_type

def buildBallpark(_session,executionIds,filter=""):
    import tempfile
    import io
    wb = openpyxl.load_workbook('ballpark_base.xlsx')
    ws = wb['Inventory']
    ws.delete_rows(2,ws.max_row)

    df_inventory = _session.table(TABLE_INPUT_FILES_INVENTORY)\
        .where(col(COLUMN_EXECUTION_ID)\
        .isin(executionIds)).select( \
        lit("").alias('"Object_Type"'),\
        col(COLUMN_ELEMENT).alias('"Schema_Name"'),\
        col(COLUMN_ELEMENT).alias('"Object_Name"'),\
        lit("").alias('"Extension"'),\
        lit("").alias('"Technology"'),\
        lit("OK").alias('"Status"'),\
        lit("FALSE").alias('"isBinary"'),\
        lit("BYTES").alias('"Bytes"'),\
        lit("CODE").alias('"ContentType"'),\
        col(COLUMN_LINES_OF_CODE).alias('"ContentLines"'),\
        lit("").alias('"CommentLines"'),\
        lit("").alias('"BlankLines"'),\
        lit("").alias('"Size"'),\
        lit("").alias('"Manual Effort HH"'),\
        lit("").alias('"Automatic Effort HH"'),\
        lit("").alias('"DF Aliases"'),\
        lit("").alias('"RDD Map"'),\
        lit("").alias('"UDF"'),\
        lit("").alias('"FS"'),\
        lit("").alias('"Added Effort"'),\
        lit("").alias('"Priority"'),\
        lit("").alias('"ID"'),\
        lit("").alias('"ConversionScope"'),\
        lit("").alias('"ConversionStatus"'),\
        lit("").alias('"ConversionDate"'),\
        lit("").alias('"Assigned to"'),\
        lit("").alias('"TestingStatus"'),\
        lit("").alias('"TestingDate"'),\
        lit("").alias('"ReTestingDate"'),\
        lit("").alias('"DeliveryDate"'),\
        lit("").alias('"Notes"'),\
        lit("").alias('"ErrorType"'),\
        lit("").alias('"ErrorMessage"')).to_pandas()

    # Apply the function to update the 'Object_Type' column
    df_inventory['Object_Type'] = df_inventory[KEY_SCHEMA_NAME].apply(getObjectType)
    df_inventory['Technology'] = df_inventory[KEY_SCHEMA_NAME].apply(getObjectType)
    df_inventory['Extension'] = df_inventory[KEY_SCHEMA_NAME].apply(getObjectType)
    df_inventory['Size'] = df_inventory[KEY_SCHEMA_NAME].apply(lambda x: size_formula)
    df_inventory['Manual Effort HH'] = df_inventory[KEY_SCHEMA_NAME].apply(lambda x: manual_time_formula)
    df_inventory['Automatic Effort HH'] = df_inventory[KEY_SCHEMA_NAME].apply(lambda x: auto_time_formula)
    df_inventory['Added Effort'] = df_inventory[KEY_SCHEMA_NAME].apply(lambda x: added_effort_formula)

    if filter:
        df_inventory = df_inventory[~df_inventory[KEY_SCHEMA_NAME].str.contains(filter)]
    for row in df_inventory.itertuples(index=False):
        ws.append([*row])
    table = ws.tables[TABLE_INVENTORY]

    from openpyxl.worksheet.table import Table
    resTable = Table(displayName=TABLE_INVENTORY,name=TABLE_INVENTORY,ref=f"A1:AG{len(df_inventory)+1}",tableColumns=table.tableColumns)
    resTable.tableStyleInfo = table.tableStyleInfo
    ws.tables[TABLE_INVENTORY] = resTable

    # Create a temporary file
    with tempfile.NamedTemporaryFile() as temp_file:
        wb.save(temp_file.name)
        # Read the contents of the temporary file into a BytesIO object
        bytes_io = io.BytesIO(temp_file.read())
        return bytes_io